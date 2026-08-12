from __future__ import annotations

import csv
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader
from pptx import Presentation

from .models import ParseResult
from .vision import (
    AsyncVisionCallback,
    VisionCallback,
    run_vision_with_fallback,
    run_vision_with_fallback_async,
)


SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".html", ".htm",
    ".doc", ".docx", ".xls", ".xlsx", ".xlsm", ".csv", ".tsv",
    ".ppt", ".pptx", ".pptm", ".pdf",
    ".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
LEGACY_TARGETS = {".doc": ".docx", ".xls": ".xlsx", ".ppt": ".pptx"}


@dataclass(slots=True)
class ParserConfig:
    max_bytes: int = 20 * 1024 * 1024
    max_rows_per_sheet: int = 500
    max_pdf_vision_pages: int = 24
    vision_prompt: str = (
        "Describe the visible objects, layout, relationships, and all readable text. "
        "Separate direct observations from uncertainty."
    )
    vision_primary: VisionCallback | None = None
    vision_fallback: VisionCallback | None = None


@dataclass(slots=True)
class AsyncParserConfig:
    max_bytes: int = 20 * 1024 * 1024
    max_rows_per_sheet: int = 500
    max_pdf_vision_pages: int = 24
    vision_prompt: str = (
        "Describe the visible objects, layout, relationships, and all readable text. "
        "Separate direct observations from uncertainty."
    )
    vision_primary: AsyncVisionCallback | None = None
    vision_fallback: AsyncVisionCallback | None = None


class AttachmentParseError(ValueError):
    pass


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _markdown_table(rows: list[list[object]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [[str(cell if cell is not None else "") for cell in row] + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
    return "\n".join(lines)


def _parse_docx(data: bytes) -> ParseResult:
    document = Document(io.BytesIO(data))
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        rendered = _markdown_table(rows)
        if rendered:
            blocks.append(rendered)
    return ParseResult("\n\n".join(blocks), {"format": "docx", "tables": len(document.tables)})


def _parse_workbook(data: bytes, max_rows: int) -> ParseResult:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True, max_row=max_rows)]
            rows = [row for row in rows if any(cell not in (None, "") for cell in row)]
            sections.append(f"## Sheet: {worksheet.title}\n\n{_markdown_table(rows)}")
    finally:
        workbook.close()
    return ParseResult("\n\n".join(sections), {"format": "xlsx", "sheets": len(sections)})


def _parse_delimited(data: bytes, delimiter: str) -> ParseResult:
    rows = list(csv.reader(io.StringIO(_decode_text(data)), delimiter=delimiter))
    return ParseResult(_markdown_table(rows), {"format": "csv" if delimiter == "," else "tsv", "rows": len(rows)})


def _parse_presentation(data: bytes) -> ParseResult:
    presentation = Presentation(io.BytesIO(data))
    sections: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                value = shape.text.strip()
                if value:
                    texts.append(value)
            if getattr(shape, "has_table", False):
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                texts.append(_markdown_table(rows))
        sections.append(f"## Slide {index}\n\n" + "\n\n".join(texts))
    return ParseResult("\n\n".join(sections), {"format": "pptx", "slides": len(sections)})


def _meaningful_length(text: str) -> int:
    return len(re.findall(r"[A-Za-z0-9\u3400-\u4dbf\u4e00-\u9fff]", text))


def _render_pdf_page(data: bytes, page_index: int) -> bytes:
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        raise AttachmentParseError("PDF image rendering requires the pdf-vision extra") from exc
    document = pdfium.PdfDocument(data)
    try:
        image = document[page_index].render(scale=2).to_pil()
        output = io.BytesIO()
        image.save(output, format="PNG")
        return output.getvalue()
    finally:
        document.close()


def _parse_pdf(data: bytes, config: ParserConfig) -> ParseResult:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and reader.decrypt("") == 0:
        return ParseResult("", {"format": "pdf", "encrypted": True}, error="encrypted_pdf")
    native_pages = [(page.extract_text() or "").strip() for page in reader.pages]
    rendered_pages = 0
    output_pages: list[str] = []
    for index, native_text in enumerate(native_pages):
        selected = native_text
        if (
            _meaningful_length(native_text) < 60
            and config.vision_primary is not None
            and rendered_pages < config.max_pdf_vision_pages
        ):
            rendered = _render_pdf_page(data, index)
            visual_text = run_vision_with_fallback(
                rendered,
                "image/png",
                config.vision_prompt,
                config.vision_primary,
                config.vision_fallback,
            )
            if visual_text:
                selected = "\n\n".join(part for part in (native_text, visual_text) if part)
            rendered_pages += 1
        if selected:
            output_pages.append(f"## Page {index + 1}\n\n{selected}")
    return ParseResult(
        "\n\n".join(output_pages),
        {"format": "pdf", "pages": len(native_pages), "vision_pages": rendered_pages},
        error="" if output_pages else "no_extractable_pdf_content",
    )


def _parse_image(data: bytes, extension: str, config: ParserConfig) -> ParseResult:
    try:
        with Image.open(io.BytesIO(data)) as image:
            image.verify()
        with Image.open(io.BytesIO(data)) as image:
            width, height = image.size
            detected_format = str(image.format or extension.lstrip(".")).lower()
    except Exception as exc:
        raise AttachmentParseError("invalid_image") from exc
    metadata = {"format": detected_format, "width": width, "height": height}
    if config.vision_primary is None:
        return ParseResult("", {**metadata, "needs_vision": True}, error="vision_required")
    mime = "image/jpeg" if extension in {".jpg", ".jpeg"} else f"image/{detected_format}"
    text = run_vision_with_fallback(
        data,
        mime,
        config.vision_prompt,
        config.vision_primary,
        config.vision_fallback,
    )
    return ParseResult(text, {**metadata, "vision_used": True}, error="" if text else "invalid_vision_output")


def _image_metadata(data: bytes, extension: str) -> tuple[dict[str, object], str]:
    try:
        with Image.open(io.BytesIO(data)) as image:
            image.verify()
        with Image.open(io.BytesIO(data)) as image:
            width, height = image.size
            detected_format = str(image.format or extension.lstrip(".")).lower()
    except Exception as exc:
        raise AttachmentParseError("invalid_image") from exc
    mime_type = "image/jpeg" if extension in {".jpg", ".jpeg"} else f"image/{detected_format}"
    return {"format": detected_format, "width": width, "height": height}, mime_type


async def _parse_image_async(
    data: bytes,
    extension: str,
    config: AsyncParserConfig,
) -> ParseResult:
    metadata, mime_type = _image_metadata(data, extension)
    if config.vision_primary is None:
        return ParseResult("", {**metadata, "needs_vision": True}, error="vision_required")
    text = await run_vision_with_fallback_async(
        data,
        mime_type,
        config.vision_prompt,
        config.vision_primary,
        config.vision_fallback,
    )
    return ParseResult(
        text,
        {**metadata, "vision_used": True},
        error="" if text else "invalid_vision_output",
    )


async def _parse_pdf_async(data: bytes, config: AsyncParserConfig) -> ParseResult:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and reader.decrypt("") == 0:
        return ParseResult("", {"format": "pdf", "encrypted": True}, error="encrypted_pdf")
    native_pages = [(page.extract_text() or "").strip() for page in reader.pages]
    rendered_pages = 0
    output_pages: list[str] = []
    for index, native_text in enumerate(native_pages):
        selected = native_text
        if (
            _meaningful_length(native_text) < 60
            and config.vision_primary is not None
            and rendered_pages < config.max_pdf_vision_pages
        ):
            rendered = _render_pdf_page(data, index)
            visual_text = await run_vision_with_fallback_async(
                rendered,
                "image/png",
                config.vision_prompt,
                config.vision_primary,
                config.vision_fallback,
            )
            if visual_text:
                selected = "\n\n".join(part for part in (native_text, visual_text) if part)
            rendered_pages += 1
        if selected:
            output_pages.append(f"## Page {index + 1}\n\n{selected}")
    return ParseResult(
        "\n\n".join(output_pages),
        {"format": "pdf", "pages": len(native_pages), "vision_pages": rendered_pages},
        error="" if output_pages else "no_extractable_pdf_content",
    )


def _convert_with_libreoffice(data: bytes, extension: str, target_extension: str) -> bytes | None:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        return None
    with tempfile.TemporaryDirectory(prefix="attachment-convert-") as directory:
        root = Path(directory)
        source = root / f"source{extension}"
        source.write_bytes(data)
        result = subprocess.run(
            [executable, "--headless", "--safe-mode", "--convert-to", target_extension.lstrip("."), "--outdir", str(root), str(source)],
            check=False,
            capture_output=True,
            timeout=120,
        )
        target = root / f"source{target_extension}"
        if result.returncode != 0 or not target.is_file():
            return None
        return target.read_bytes()


def _convert_with_windows_office(data: bytes, extension: str, target_extension: str) -> bytes | None:
    if os.name != "nt":
        return None
    try:
        import win32com.client
    except ImportError:
        return None
    with tempfile.TemporaryDirectory(prefix="attachment-convert-") as directory:
        root = Path(directory)
        source = root / f"source{extension}"
        target = root / f"source{target_extension}"
        source.write_bytes(data)
        creation_flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        result = subprocess.run(
            [sys.executable, "-m", "langbot_unified_pipeline.office_worker", str(source), str(target), extension],
            check=False,
            capture_output=True,
            timeout=120,
            creationflags=creation_flags,
        )
        if result.returncode != 0 or not target.is_file():
            return None
        return target.read_bytes()


def _convert_legacy(data: bytes, extension: str) -> tuple[bytes, str]:
    target = LEGACY_TARGETS[extension]
    converted = _convert_with_libreoffice(data, extension, target)
    if converted is None:
        converted = _convert_with_windows_office(data, extension, target)
    if converted is None:
        raise AttachmentParseError("legacy_office_converter_unavailable")
    return converted, target


def parse_attachment(data: bytes, filename: str, config: ParserConfig | None = None) -> ParseResult:
    config = config or ParserConfig()
    extension = Path(filename).suffix.casefold()
    if extension not in SUPPORTED_EXTENSIONS:
        raise AttachmentParseError(f"unsupported_extension:{extension or '<none>'}")
    if not data or len(data) > config.max_bytes:
        raise AttachmentParseError("empty_or_oversized_payload")
    if extension in LEGACY_TARGETS:
        data, extension = _convert_legacy(data, extension)
    try:
        if extension in {".txt", ".md", ".markdown"}:
            return ParseResult(_decode_text(data).strip(), {"format": extension.lstrip(".")})
        if extension in {".html", ".htm"}:
            soup = BeautifulSoup(_decode_text(data), "html.parser")
            for element in soup(["script", "style", "noscript"]):
                element.decompose()
            return ParseResult(soup.get_text("\n", strip=True), {"format": "html"})
        if extension == ".docx":
            return _parse_docx(data)
        if extension in {".xlsx", ".xlsm"}:
            return _parse_workbook(data, config.max_rows_per_sheet)
        if extension in {".csv", ".tsv"}:
            return _parse_delimited(data, "," if extension == ".csv" else "\t")
        if extension in {".pptx", ".pptm"}:
            return _parse_presentation(data)
        if extension == ".pdf":
            return _parse_pdf(data, config)
        if extension in IMAGE_EXTENSIONS:
            return _parse_image(data, extension, config)
    except AttachmentParseError:
        raise
    except Exception as exc:
        raise AttachmentParseError(f"parse_failed:{type(exc).__name__}") from exc
    raise AttachmentParseError(f"unsupported_extension:{extension}")


async def parse_attachment_async(
    data: bytes,
    filename: str,
    config: AsyncParserConfig | None = None,
) -> ParseResult:
    config = config or AsyncParserConfig()
    extension = Path(filename).suffix.casefold()
    if extension not in SUPPORTED_EXTENSIONS:
        raise AttachmentParseError(f"unsupported_extension:{extension or '<none>'}")
    if not data or len(data) > config.max_bytes:
        raise AttachmentParseError("empty_or_oversized_payload")
    if extension in LEGACY_TARGETS:
        data, extension = _convert_legacy(data, extension)
        filename = f"{Path(filename).stem}{extension}"
    try:
        if extension == ".pdf":
            return await _parse_pdf_async(data, config)
        if extension in IMAGE_EXTENSIONS:
            return await _parse_image_async(data, extension, config)
        return parse_attachment(
            data,
            filename,
            ParserConfig(
                max_bytes=config.max_bytes,
                max_rows_per_sheet=config.max_rows_per_sheet,
                max_pdf_vision_pages=config.max_pdf_vision_pages,
                vision_prompt=config.vision_prompt,
            ),
        )
    except AttachmentParseError:
        raise
    except Exception as exc:
        raise AttachmentParseError(f"parse_failed:{type(exc).__name__}") from exc
