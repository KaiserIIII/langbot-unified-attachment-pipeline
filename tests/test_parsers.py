from __future__ import annotations

import io

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from reportlab.pdfgen.canvas import Canvas

from langbot_unified_pipeline.parsers import AttachmentParseError, ParserConfig, parse_attachment
from langbot_unified_pipeline import parsers


MARKER = "PUBLIC_TEST_ALPHA"


def _docx() -> bytes:
    output = io.BytesIO()
    document = Document()
    document.add_heading(MARKER, level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "field"
    table.cell(0, 1).text = "value"
    table.cell(1, 0).text = "status"
    table.cell(1, 1).text = "verified"
    document.save(output)
    return output.getvalue()


def _xlsx() -> bytes:
    output = io.BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["marker", "value"])
    worksheet.append([MARKER, 42])
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _pptx() -> bytes:
    output = io.BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = MARKER
    slide.placeholders[1].text = "Presentation extraction fixture"
    presentation.save(output)
    return output.getvalue()


def _pdf() -> bytes:
    output = io.BytesIO()
    canvas = Canvas(output)
    canvas.drawString(72, 720, MARKER)
    canvas.drawString(72, 700, "Native PDF text extraction fixture")
    canvas.save()
    return output.getvalue()


def _blank_pdf() -> bytes:
    output = io.BytesIO()
    canvas = Canvas(output)
    canvas.showPage()
    canvas.save()
    return output.getvalue()


def _image(image_format: str = "PNG") -> bytes:
    output = io.BytesIO()
    image = Image.new("RGB", (120, 80), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((30, 20, 90, 60), fill="red")
    image.save(output, format=image_format)
    return output.getvalue()


@pytest.mark.parametrize(
    ("filename", "payload", "expected"),
    [
        ("sample.txt", MARKER.encode(), MARKER),
        ("sample.md", f"# {MARKER}".encode(), MARKER),
        ("sample.html", f"<html><script>ignore()</script><h1>{MARKER}</h1></html>".encode(), MARKER),
        ("sample.csv", f"marker,value\n{MARKER},42\n".encode(), MARKER),
        ("sample.tsv", f"marker\tvalue\n{MARKER}\t42\n".encode(), MARKER),
        ("sample.docx", _docx(), MARKER),
        ("sample.xlsx", _xlsx(), MARKER),
        ("sample.xlsm", _xlsx(), MARKER),
        ("sample.pdf", _pdf(), MARKER),
        ("sample.pptx", _pptx(), MARKER),
        ("sample.pptm", _pptx(), MARKER),
    ],
    ids=["txt", "markdown", "html", "csv", "tsv", "docx", "xlsx", "xlsm", "pdf", "pptx", "pptm"],
)
def test_common_format_matrix(filename, payload, expected):
    result = parse_attachment(payload, filename)
    assert result.ok
    assert expected in result.text


@pytest.mark.parametrize(
    ("filename", "image_format"),
    [
        ("sample.png", "PNG"),
        ("sample.jpg", "JPEG"),
        ("sample.jpeg", "JPEG"),
        ("sample.webp", "WEBP"),
        ("sample.gif", "GIF"),
        ("sample.bmp", "BMP"),
        ("sample.tif", "TIFF"),
        ("sample.tiff", "TIFF"),
    ],
)
def test_image_formats_use_semantic_vision_callback(filename, image_format):
    result = parse_attachment(
        _image(image_format),
        filename,
        ParserConfig(vision_primary=lambda _data, _mime, _prompt: f"A red rectangle. {MARKER}"),
    )
    assert result.ok
    assert "red rectangle" in result.text
    assert result.metadata["vision_used"] is True


def test_invalid_image_fails_closed():
    with pytest.raises(AttachmentParseError, match="invalid_image"):
        parse_attachment(b"not an image", "sample.png")


def test_image_without_vision_callback_requires_vision():
    result = parse_attachment(_image(), "sample.png")
    assert not result.ok
    assert result.error == "vision_required"
    assert result.metadata["needs_vision"] is True


def test_blank_pdf_uses_bounded_page_vision():
    result = parse_attachment(
        _blank_pdf(),
        "scan.pdf",
        ParserConfig(vision_primary=lambda _data, _mime, _prompt: f"Scanned marker {MARKER}"),
    )
    assert result.ok
    assert MARKER in result.text
    assert result.metadata["vision_pages"] == 1


def test_blank_pdf_without_vision_fails_closed():
    result = parse_attachment(_blank_pdf(), "blank.pdf")
    assert not result.ok
    assert result.error == "no_extractable_pdf_content"


@pytest.mark.parametrize(
    ("legacy_name", "modern_payload", "expected"),
    [("sample.doc", _docx(), MARKER), ("sample.xls", _xlsx(), MARKER), ("sample.ppt", _pptx(), MARKER)],
    ids=["doc", "xls", "ppt"],
)
def test_legacy_office_routes_converted_bytes(monkeypatch, legacy_name, modern_payload, expected):
    monkeypatch.setattr(parsers, "_convert_with_libreoffice", lambda _data, _source, _target: modern_payload)
    result = parse_attachment(b"legacy payload", legacy_name)
    assert result.ok
    assert expected in result.text


def test_legacy_office_without_converter_fails_closed(monkeypatch):
    monkeypatch.setattr(parsers, "_convert_with_libreoffice", lambda *_args: None)
    monkeypatch.setattr(parsers, "_convert_with_windows_office", lambda *_args: None)
    with pytest.raises(AttachmentParseError, match="legacy_office_converter_unavailable"):
        parse_attachment(b"legacy payload", "sample.doc")


def test_oversized_payload_is_rejected():
    with pytest.raises(AttachmentParseError, match="empty_or_oversized_payload"):
        parse_attachment(b"1234", "sample.txt", ParserConfig(max_bytes=3))


def test_unsupported_extension_fails_closed():
    with pytest.raises(AttachmentParseError, match="unsupported_extension"):
        parse_attachment(b"payload", "sample.exe")
